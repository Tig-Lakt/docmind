import httpx
from langchain_text_splitters import RecursiveCharacterTextSplitter


def extract_text(file_path: str, file_type: str) -> str:
    if file_type == "pdf":
        return _extract_pdf(file_path)
    elif file_type == "docx":
        return _extract_docx(file_path)
    elif file_type == "xlsx":
        return _extract_xlsx(file_path)
    raise ValueError(f"Неподдерживаемый тип: {file_type}")


def _extract_pdf(path: str) -> str:
    import pdfplumber
    text = []
    with pdfplumber.open(path) as pdf:
        for page in pdf.pages:
            text.append(page.extract_text() or "")
    return "\n".join(text)


def _extract_docx(path: str) -> str:
    from docx import Document
    doc = Document(path)
    return "\n".join(p.text for p in doc.paragraphs if p.text.strip())


def _extract_xlsx(path: str) -> str:
    import openpyxl
    wb = openpyxl.load_workbook(path, data_only=True)
    lines = []
    for sheet in wb.worksheets:
        for row in sheet.iter_rows(values_only=True):
            line = " | ".join(str(c) for c in row if c is not None)
            if line:
                lines.append(line)
    return "\n".join(lines)


async def extract_text_from_url(url: str) -> str:
    async with httpx.AsyncClient() as client:
        response = await client.get(url, timeout=15)
        response.raise_for_status()
        return response.text


def split_text(text: str) -> list[str]:
    splitter = RecursiveCharacterTextSplitter(
        chunk_size=500,
        chunk_overlap=50,
        separators=["\n\n", "\n", ".", " "],
    )
    return splitter.split_text(text)